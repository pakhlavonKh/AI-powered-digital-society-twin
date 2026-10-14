"""Extract and validate Ankara 2025 age/sex counts from the TUIK T16 CSV."""

from __future__ import annotations

import argparse
import csv
import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
# This workspace is itself the project's data directory.
RAW_DIR = PROJECT_ROOT / "raw" / "tuik"
OUTPUT_PATH = PROJECT_ROOT / "processed" / "ankara_age_gender.csv"
CONTROLS_PATH = PROJECT_ROOT / "processed" / "ankara_demographic_controls.csv"
AGE_GROUPS = tuple(f"{age}-{age + 4}" for age in range(0, 90, 5)) + ("90+",)
EXPECTED_TOTALS = {"male": 2911853, "female": 2998467, "total": 5910320}


@dataclass
class CandidateSheet:
    path: Path
    sheet_name: str
    rows: list[list[object]]


def normalize(value):
    text = str(value).strip().lower().replace("\u0131", "i")
    return " ".join("".join(c for c in unicodedata.normalize("NFKD", text)
                            if not unicodedata.combining(c)).replace("_", " ").split())


def iter_workbook_sheets(raw_dir):
    for path in sorted(raw_dir.rglob("*")):
        if path.name.startswith("~$") or path.suffix.lower() not in {".xls", ".xlsx"}:
            continue
        if path.suffix.lower() == ".xls":
            import xlrd
            book = xlrd.open_workbook(str(path))
            for sheet in book.sheets():
                yield CandidateSheet(path, sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)])
        else:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    yield CandidateSheet(path, sheet.title, [list(r) for r in sheet.iter_rows(values_only=True)])
            finally:
                book.close()


def to_number(value):
    # Accept Excel numbers and explicit grouped integer strings, never truncate fractions.
    if isinstance(value, str):
        value = value.strip()
        if re.fullmatch(r"\d{1,3}(?:[., ]\d{3})+", value):
            value = re.sub(r"[., ]", "", value)
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Invalid population: {value!r}")
    number = float(value)
    if not math.isfinite(number) or number < 0 or not number.is_integer():
        raise ValueError(f"Population must be a nonnegative integer: {value!r}")
    return int(number)


def year_value(value):
    year = to_number(value)
    if not 1900 <= year <= 2100:
        raise ValueError(f"Invalid year: {value!r}")
    return year


def title_year(rows):
    years = set(re.findall(r"\b(?:19|20)\d{2}\b", " ".join(str(v) for r in rows for v in r)))
    if len(years) != 1:
        raise ValueError("A year column or a single unambiguous title year is required")
    return int(years.pop())


ALIASES = {
    "year": {"year", "yil"},
    "province": {"province", "provinces", "il", "il provinces", "il province"},
    "age_group": {"age group", "yas grubu", "yas grubu age group"},
    "gender": {"gender", "sex", "cinsiyet", "cinsiyet sex"},
    "population": {"population", "nufus", "nufus population"},
    "male": {"male", "erkek", "erkek male"},
    "female": {"female", "kadin", "kadin female"},
}


def extract_tidy_rows(sheet):
    for index, header in enumerate(sheet.rows):
        columns = {}
        for key, aliases in ALIASES.items():
            matches = [i for i, value in enumerate(header) if normalize(value) in aliases]
            if len(matches) == 1:
                columns[key] = matches[0]
        tidy = {"gender", "population"} <= columns.keys()
        wide = {"male", "female"} <= columns.keys()
        if not {"province", "age_group"} <= columns.keys() or not (tidy or wide):
            continue
        results = []
        province = ""
        for row in sheet.rows[index + 1:]:
            def get(key):
                pos = columns[key]
                return row[pos] if pos < len(row) else None
            # Merged province cells can extend across consecutive age rows.
            if get("province") not in (None, ""):
                province = normalize(get("province"))
            if province != "ankara":
                continue
            age = normalize(get("age_group"))
            if age in {"total", "toplam", "toplam-total", "toplam total"}:
                continue
            match = re.fullmatch(r"(\d{1,3})\s*[-\u2013]\s*(\d{1,3})|((?:\d{1,3}))\s*\+", age)
            if not match:
                if age in {"", "none"} and all(v in (None, "") for v in row):
                    province = ""
                    continue
                raise ValueError(f"Unrecognized Ankara age group {age!r} in {sheet.path.name}")
            age = f"{match[1]}-{match[2]}" if match[1] else f"{match[3]}+"
            year = year_value(get("year")) if "year" in columns else title_year(sheet.rows[:index])
            pairs = [(normalize(get("gender")), get("population"))] if tidy else [(g, get(g)) for g in ("male", "female")]
            for gender, population in pairs:
                gender = {"erkek": "male", "kadin": "female"}.get(gender, gender)
                results.append(dict(year=year, province="Ankara", age_group=age,
                                    gender=gender, population=to_number(population)))
        return results
    return []


def validate_rows(rows):
    if not rows:
        raise ValueError("No Ankara age-by-gender observations found")
    unique = {}
    for row in rows:
        if row["gender"] not in {"male", "female"} or not row["age_group"]:
            raise ValueError("Missing or invalid age/gender")
        key = tuple(row[k] for k in ("year", "province", "age_group", "gender"))
        if key in unique and unique[key]["population"] != row["population"]:
            raise ValueError(f"Conflicting observations: {key}")
        unique[key] = row.copy()
    rows = list(unique.values())
    totals = {}
    groups = {}
    for row in rows:
        year = row["year"]
        row["population"] = to_number(row["population"])
        totals[year] = totals.get(year, 0) + row["population"]
        groups.setdefault((year, row["age_group"]), set()).add(row["gender"])
    if any(g != {"male", "female"} for g in groups.values()) or any(t <= 0 for t in totals.values()):
        raise ValueError("Each age group needs both genders and each year needs positive population")
    for year in totals:
        ages = sorted((age for y, age in groups if y == year), key=lambda a: int(re.match(r"\d+", a)[0]))
        next_age = 0
        for index, age in enumerate(ages):
            numbers = [int(n) for n in re.findall(r"\d+", age)]
            if numbers[0] != next_age:
                raise ValueError(f"Age groups have gaps or overlap in {year}: {age}")
            if age.endswith("+"):
                if index != len(ages) - 1:
                    raise ValueError("Open age group must be last")
            else:
                if numbers[1] < numbers[0]:
                    raise ValueError(f"Invalid age range: {age}")
                next_age = numbers[1] + 1
        if not ages[-1].endswith("+"):
            raise ValueError("Complete coverage requires a final open age group")
    for row in rows:
        row["probability"] = row["population"] / totals[row["year"]]
    for year in totals:
        if not math.isclose(sum(r["probability"] for r in rows if r["year"] == year), 1, abs_tol=1e-9):
            raise ValueError("Probabilities do not sum to one")
    return sorted(rows, key=lambda r: (r["year"], int(re.match(r"\d+", r["age_group"])[0]), r["gender"]))


def write_csv(rows, output_path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def extract_t16_csv(path):
    expected_headers = ["cografi kapsam", "gozlem sikligi", "istatistiksel gosterge",
                        "ikamet edilen yer", "cinsiyet", "yas grubu", "adnks gosterge",
                        "zaman", "gozlem", "gozlem durumu", "zaman formati",
                        "ondalik basamak sayisi", "olcu birimi", "birim carpani"]
    rows, aggregates = [], {}
    seen = set()
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=";")
        if [normalize(h) for h in (reader.fieldnames or [])] != expected_headers:
            raise ValueError("CSV headers do not match the inspected T16 format")
        for line, raw in enumerate(reader, start=2):
            if None in raw or any(v is None for v in raw.values()):
                raise ValueError(f"Malformed CSV row {line}")
            record = dict(zip(expected_headers, raw.values()))
            if record["zaman"] != "2025" or record["ikamet edilen yer"] != "Ankara":
                continue
            if normalize(record["gozlem sikligi"]) != "yillik" or normalize(record["adnks gosterge"]) != "nufus":
                raise ValueError(f"Unexpected frequency/measure on row {line}")
            if any(record[k] for k in ("gozlem durumu", "birim carpani", "olcu birimi")):
                raise ValueError(f"Unrecognized observation metadata on row {line}")
            gender = {"erkek": "male", "kadin": "female", "toplam": "total"}.get(normalize(record["cinsiyet"]))
            age = record["yas grubu"].strip()
            if gender is None or age not in (*AGE_GROUPS, "Toplam"):
                raise ValueError(f"Unexpected gender or age on row {line}")
            key = (age, gender)
            if key in seen:
                raise ValueError(f"Duplicate Ankara 2025 observation: {key}")
            seen.add(key)
            if not re.fullmatch(r"\d+", record["gozlem"]):
                raise ValueError(f"Expected integer population on row {line}")
            population = int(record["gozlem"])
            if age == "Toplam" or gender == "total":
                aggregates[key] = population
            else:
                rows.append(dict(year=2025, province="Ankara", age_group=age,
                                 gender=gender, population=population))
    validated = validate_ankara_2025(rows)
    for (age, gender), population in aggregates.items():
        actual = sum(r["population"] for r in validated
                     if (age == "Toplam" or r["age_group"] == age)
                     and (gender == "total" or r["gender"] == gender))
        if actual != population:
            raise ValueError(f"Source aggregate mismatch for {(age, gender)}")
    return validated


def validate_ankara_2025(rows):
    if len(rows) != 38:
        raise ValueError(f"Expected exactly 38 observations, found {len(rows)}")
    expected = {(2025, "Ankara", age, gender) for age in AGE_GROUPS for gender in ("male", "female")}
    keys = [tuple(r[k] for k in ("year", "province", "age_group", "gender")) for r in rows]
    if len(set(keys)) != 38 or set(keys) != expected:
        raise ValueError("Missing/duplicate age-gender cells or wrong year/province")
    validated = validate_rows(rows)
    for gender, expected_total in EXPECTED_TOTALS.items():
        actual = sum(r["population"] for r in validated if gender == "total" or r["gender"] == gender)
        if actual != expected_total:
            raise ValueError(f"{gender} population {actual} != {expected_total}")
    return validated


def reconcile_controls(rows, controls_path):
    controls = {}
    with controls_path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["year"] == "2025" and row["province"] == "Ankara" and row["statistic"] in {
                "male_population", "female_population", "total_population"
            }:
                value = to_number(row["value"])
                if row["statistic"] in controls and controls[row["statistic"]] != value:
                    raise ValueError("Conflicting population controls")
                controls[row["statistic"]] = value
    for gender in EXPECTED_TOTALS:
        actual = sum(r["population"] for r in rows if gender == "total" or r["gender"] == gender)
        if controls.get(f"{gender}_population") != actual:
            raise ValueError(f"Missing or mismatched {gender} population control")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=RAW_DIR)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    parser.add_argument("--source", type=Path, help="Explicit T16 CSV; otherwise locate it under --raw-dir")
    parser.add_argument("--controls", type=Path, default=CONTROLS_PATH)
    args = parser.parse_args()
    candidates = [args.source] if args.source else sorted(args.raw_dir.rglob("*DF_ADNKS_T16,1.0*.csv"))
    if len(candidates) != 1:
        parser.error(f"Expected one T16 CSV; found {len(candidates)}. Use --source to disambiguate.")
    validated = extract_t16_csv(candidates[0])
    reconcile_controls(validated, args.controls)
    write_csv(validated, args.output)
    print(f"Source: {candidates[0]}")
    print(f"PASS: 38 cells, 19 age groups, both genders, no duplicates; totals {EXPECTED_TOTALS}; controls match")
    print(f"Probability sum: {math.fsum(r['probability'] for r in validated):.12f}")
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
